from __future__ import annotations

import math
import re
import threading
import unicodedata
from collections import Counter, defaultdict
import hashlib
import json
import time
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Callable, Iterable

from openpyxl import load_workbook


REPORT_HEADERS = [
    "Строка источника",
    "Статус",
    "№ с/з",
    "№ с/р",
    "Дата с/р",
    "Получатель",
    "Номенклатура по с/р",
    "Наименование",
    "Код номенклатуры",
    "Ед. изм.",
    "Потребность",
    "Закуплено",
    "Руководитель",
    "Исполнитель",
    "Ответственный (S)",
    "Срок ответственного ОМТС",
    "Комментарий",
    "Объект",
    "ЦФО",
    "ОКПД2",
    "Наименование ОКПД2",
    "Проверка ОКПД2",
]

STOPWORDS = {
    "а",
    "без",
    "в",
    "для",
    "до",
    "и",
    "из",
    "к",
    "на",
    "над",
    "не",
    "о",
    "об",
    "от",
    "по",
    "под",
    "при",
    "с",
    "у",
    "через",
}


def normalize_text(value: object) -> str:
    text = unicodedata.normalize("NFC", "" if value is None else str(value)).replace("ё", "е").replace("Ё", "Е")
    return re.sub(r"\s+", " ", text).strip()


def person_key(value: object) -> str:
    return re.sub(r"[\s.]", "", normalize_text(value).casefold())


def split_responsible(value: object) -> list[str]:
    return [part for part in (normalize_text(item) for item in str(value or "").split("/")) if part]


def display_value(value: object, identifier: bool = False) -> object:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value)) if identifier else int(value)
    if identifier:
        return normalize_text(value)
    return value


def tokenize(value: object) -> list[str]:
    return [
        token
        for token in re.findall(r"[a-zа-я0-9]+", normalize_text(value).casefold())
        if len(token) > 1 and not token.isdigit() and token not in STOPWORDS
    ]


def positions_from_rows(rows: Iterable[Iterable[object]], first_row: int = 2) -> list[dict[str, object]]:
    """Позиции «В работе» из строк листа заявок (A — статус, S — ответственные через «/»).

    Одинаково для Excel-выгрузки и живой Google-таблицы: Google API не возвращает
    пустые ячейки в конце строки, поэтому строки дополняются до 29 столбцов (A..AC).
    """
    positions = []
    for row_number, row in enumerate(rows, start=first_row):
        row = tuple(row)
        if len(row) < 29:
            row = row + (None,) * (29 - len(row))
        status = normalize_text(row[0])
        if status.casefold() != "в работе":
            continue
        responsible = normalize_text(row[18])
        people = split_responsible(responsible)
        item_name = normalize_text(row[8])
        if not responsible and not item_name:
            continue
        positions.append(
            {
                "source_row": row_number,
                "status": status,
                "request_number": display_value(row[3], identifier=True),
                "service_request_number": display_value(row[4], identifier=True),
                "service_request_date": display_value(row[5]),
                "recipient": normalize_text(row[6]),
                "source_nomenclature": display_value(row[7], identifier=True),
                "item_name": item_name,
                "item_code": display_value(row[14], identifier=True),
                "unit": normalize_text(row[9]),
                "quantity": display_value(row[10]),
                "purchased": display_value(row[12]),
                "manager": people[0] if people else "",
                "employee": " / ".join(people[1:]) if len(people) > 1 else (people[0] if people else ""),
                "responsible": responsible,
                "deadline": display_value(row[20]),
                "comment": normalize_text(row[21]),
                "object": normalize_text(row[27]),
                "cost_center": normalize_text(row[28]),
            }
        )
    return positions


class WorkReport:
    def __init__(self, workbook_path: Path, okpd_path: Path, sheet_name: str = "незакр 202526") -> None:
        self.workbook_path = Path(workbook_path)
        self.okpd_path = Path(okpd_path)
        self.sheet_name = sheet_name
        self._lock = threading.RLock()
        self._workbook_stamp: tuple[int, int] | None = None
        self._positions: list[dict[str, object]] = []
        self._okpd_stamp: tuple[int, int] | None = None
        self._okpd_rows: list[tuple[str, str]] = []
        self._okpd_by_code: dict[str, tuple[str, str]] = {}
        self._documents: list[list[str]] = []
        self._document_frequency: Counter[str] = Counter()
        self._postings: dict[str, list[tuple[int, int]]] = {}
        self._vocabulary: set[str] = set()
        self._vocabulary_list: list[str] = []

    @staticmethod
    def _stamp(path: Path) -> tuple[int, int]:
        stat = path.stat()
        return stat.st_mtime_ns, stat.st_size

    @property
    def source_name(self) -> str:
        return self.workbook_path.name

    def source_version(self) -> str:
        stat = self.workbook_path.stat()
        return f"{stat.st_mtime_ns}:{stat.st_size}"

    def source_modified(self) -> datetime | None:
        return datetime.fromtimestamp(self.workbook_path.stat().st_mtime, timezone.utc)

    def _load_positions(self) -> None:
        stamp = self._stamp(self.workbook_path)
        workbook = load_workbook(self.workbook_path, read_only=True, data_only=True, keep_links=False)
        try:
            if self.sheet_name not in workbook.sheetnames:
                raise ValueError(f"Лист '{self.sheet_name}' не найден в {self.workbook_path.name}")
            positions = positions_from_rows(workbook[self.sheet_name].iter_rows(min_row=2, values_only=True))
            if stamp != self._stamp(self.workbook_path):
                raise ValueError("Файл заявок изменился во время чтения; повторите запрос")
            self._positions = positions
            self._workbook_stamp = stamp
        finally:
            workbook.close()

    def _load_okpd(self) -> None:
        stamp = self._stamp(self.okpd_path)
        workbook = load_workbook(self.okpd_path, read_only=True, data_only=True, keep_links=False)
        try:
            sheet = workbook["Лист1"]
            rows = []
            documents = []
            frequency: Counter[str] = Counter()
            postings: defaultdict[str, list[tuple[int, int]]] = defaultdict(list)
            by_code = {}
            for row in sheet.iter_rows(min_row=8, min_col=2, max_col=3, values_only=True):
                code = re.sub(r"\s+", "", normalize_text(row[0]))
                title = normalize_text(row[1])
                # Letter section headings (A..U) are not assignable OKPD2 codes.
                if not re.fullmatch(r"\d{2}(?:\.\d{1,3}){0,3}", code) or not title:
                    continue
                index = len(rows)
                rows.append((code, title))
                by_code.setdefault(code.casefold(), (code, title))
                tokens = tokenize(title)
                documents.append(tokens)
                term_frequencies = Counter(tokens)
                frequency.update(term_frequencies.keys())
                for token, term_frequency in term_frequencies.items():
                    postings[token].append((index, term_frequency))
            if not rows:
                raise ValueError("Справочник ОКПД2 пуст")
            if stamp != self._stamp(self.okpd_path):
                raise ValueError("Классификатор изменился во время чтения; повторите запрос")
            self._okpd_rows = rows
            self._okpd_by_code = by_code
            self._documents = documents
            self._document_frequency = frequency
            self._postings = dict(postings)
            self._vocabulary = set(frequency)
            self._vocabulary_list = sorted(self._vocabulary)
            self._okpd_stamp = stamp
        finally:
            workbook.close()

    def _ensure_positions(self) -> None:
        stamp = self._stamp(self.workbook_path)
        with self._lock:
            if self._workbook_stamp is None or self._workbook_stamp != stamp:
                self._load_positions()

    def _ensure_okpd(self) -> None:
        stamp = self._stamp(self.okpd_path)
        with self._lock:
            if self._okpd_stamp != stamp:
                self._load_okpd()

    def positions(self, employee: str, manager: str | None = None) -> list[dict[str, object]]:
        self._ensure_positions()
        query_people = split_responsible(employee)
        if not query_people:
            raise ValueError("Укажите сотрудника или пару 'Руководитель/Сотрудник'")
        query_keys = [person_key(person) for person in query_people]
        manager_key = person_key(manager)
        if len(query_keys) > 1 and manager_key and manager_key != query_keys[0]:
            raise ValueError("Руководитель в параметрах не совпадает с парой 'Руководитель/Сотрудник'")
        result = []
        for position in self._positions:
            people = split_responsible(position["responsible"])
            keys = [person_key(person) for person in people]
            if manager_key and (not keys or keys[0] != manager_key):
                continue
            if len(query_keys) > 1:
                if keys != query_keys:
                    continue
            elif query_keys[0] not in keys:
                continue
            result.append(dict(position))
        return result

    def okpd_lookup(self, code: str) -> tuple[str, str]:
        self._ensure_okpd()
        key = re.sub(r"\s+", "", normalize_text(code)).casefold()
        try:
            return self._okpd_by_code[key]
        except KeyError as exc:
            raise ValueError(f"Код ОКПД2 '{code}' отсутствует в приложенном классификаторе") from exc

    def search_okpd2(self, query: str, limit: int = 5) -> list[dict[str, object]]:
        self._ensure_okpd()
        if not normalize_text(query):
            raise ValueError("Укажите наименование для поиска ОКПД2")
        if not 1 <= limit <= 20:
            raise ValueError("limit должен быть от 1 до 20")
        code_query = re.sub(r"\s+", "", query)
        if re.fullmatch(r"\d{2}(?:\.\d{1,3}){0,3}", code_query):
            matches = [(code, title) for code, title in self._okpd_rows
                       if code == code_query or code.startswith(code_query + ".")]
            return [{"code": code, "title": title, "score": None}
                    for code, title in matches[:limit]]
        query_tokens = list(dict.fromkeys(tokenize(query)))
        if not query_tokens:
            raise ValueError("В запросе нет значимых слов для поиска ОКПД2")
        import difflib

        scores: defaultdict[int, float] = defaultdict(float)
        matched_query_tokens: defaultdict[int, set[str]] = defaultdict(set)
        total = len(self._okpd_rows)
        for query_token in query_tokens:
            related = [(query_token, 1.0)]
            if query_token not in self._vocabulary:
                if not re.fullmatch(r"[а-я]{4,}", query_token):
                    continue
                related = []
                for token in difflib.get_close_matches(query_token, self._vocabulary_list, n=5, cutoff=0.82):
                    similarity = difflib.SequenceMatcher(None, query_token, token).ratio()
                    related.append((token, similarity**3))
            for token, similarity_weight in related:
                frequency = self._document_frequency.get(token, 0)
                if not frequency:
                    continue
                inverse_frequency = math.log(1 + (total - frequency + 0.5) / (frequency + 0.5))
                for index, term_frequency in self._postings.get(token, []):
                    length_factor = 1 + 0.15 * math.log1p(len(self._documents[index]))
                    scores[index] += inverse_frequency * (1 + math.log(term_frequency)) * similarity_weight / length_factor
                    matched_query_tokens[index].add(query_token)
        normalized_query = normalize_text(query).casefold()
        ranked = []
        for index, score in scores.items():
            code, title = self._okpd_rows[index]
            coverage = len(matched_query_tokens[index]) / len(query_tokens)
            score *= 0.5 + 0.5 * coverage
            if code.count(".") >= 2:
                score *= 1.05
            if normalized_query and normalized_query in title.casefold():
                score *= 1.2
            ranked.append((score, index, code, title))
        ranked.sort(key=lambda item: (-item[0], item[2]))
        return [
            {"code": code, "title": title, "score": round(score, 4)}
            for score, _, code, title in ranked[:limit]
            if score > 0
        ]

    def build_rows(
        self,
        employee: str,
        assignments: dict[int, str],
        manager: str | None = None,
    ) -> tuple[list[dict[str, object]], list[list[object]]]:
        positions = self.positions(employee, manager)
        if not positions:
            raise ValueError("Позиции 'В работе' для указанного сотрудника не найдены")
        required = {int(position["source_row"]) for position in positions}
        selected = {int(row): code for row, code in assignments.items()}
        missing = sorted(required - set(selected))
        extra = sorted(set(selected) - required)
        if missing or extra:
            parts = []
            if missing:
                parts.append(f"не выбран ОКПД2 для строк: {missing}")
            if extra:
                parts.append(f"строки не относятся к выборке: {extra}")
            raise ValueError("; ".join(parts))
        rows = [REPORT_HEADERS.copy()]
        for position in positions:
            code, title = self.okpd_lookup(selected[int(position["source_row"])])
            rows.append(
                [
                    position["source_row"],
                    position["status"],
                    position["request_number"],
                    position["service_request_number"],
                    position["service_request_date"],
                    position["recipient"],
                    position["source_nomenclature"],
                    position["item_name"],
                    position["item_code"],
                    position["unit"],
                    position["quantity"],
                    position["purchased"],
                    position["manager"],
                    position["employee"],
                    position["responsible"],
                    position["deadline"],
                    position["comment"],
                    position["object"],
                    position["cost_center"],
                    code,
                    title,
                    "Подобрано по классификатору; требуется проверка специалиста",
                ]
            )
        return positions, rows


class GoogleSheetReport(WorkReport):
    """Заявки из живой Google-таблицы вместо Excel-выгрузки; ОКПД2 — из локального файла.

    Строки читаются под Google-аккаунтом сотрудника и кешируются на CACHE_SECONDS.
    source_version зависит только от позиций «В работе» (строка, наименование,
    ответственные): правки комментариев не ломают create_work_report, а вставка
    строк выше — ломает, как и должна, ведь сдвигаются номера source_row.
    """

    CACHE_SECONDS = 60

    def __init__(self, spreadsheet_id: str, sheet_name: str, okpd_path: Path,
                 sheets_service: Callable[[], object]) -> None:
        super().__init__(Path(f"google-sheet-{spreadsheet_id}"), okpd_path, sheet_name)
        self.spreadsheet_id = spreadsheet_id
        self._sheets_service = sheets_service
        self._fetched_at: float | None = None
        self._fetched_utc: datetime | None = None
        self._version = ""

    @property
    def source_name(self) -> str:
        return f"Google-таблица {self.spreadsheet_id}, лист «{self.sheet_name}»"

    def source_version(self) -> str:
        self._ensure_positions()
        return self._version

    def source_modified(self) -> datetime | None:
        self._ensure_positions()
        return self._fetched_utc

    def _ensure_positions(self) -> None:
        with self._lock:
            if self._fetched_at is not None and time.monotonic() - self._fetched_at < self.CACHE_SECONDS:
                return
            sheet = self.sheet_name.replace("'", "''")
            values = self._sheets_service().spreadsheets().values().get(
                spreadsheetId=self.spreadsheet_id, range=f"'{sheet}'!A2:AC",
                valueRenderOption="FORMATTED_VALUE", dateTimeRenderOption="FORMATTED_STRING",
            ).execute().get("values", [])
            positions = positions_from_rows(values)
            key = [(item["source_row"], item["item_name"], item["responsible"]) for item in positions]
            self._version = hashlib.sha256(json.dumps(key, ensure_ascii=False).encode("utf-8")).hexdigest()[:16]
            self._positions = positions
            self._fetched_at = time.monotonic()
            self._fetched_utc = datetime.now(timezone.utc)


def assignments_from_items(items: Iterable[dict[str, object]]) -> dict[int, str]:
    result = {}
    for item in items:
        if not isinstance(item, dict) or "source_row" not in item or "okpd2" not in item:
            raise ValueError("Каждое назначение должно содержать source_row и okpd2")
        row = item["source_row"]
        if type(row) is not int or row < 2:
            raise ValueError("source_row должен быть целым номером строки Excel от 2")
        if not isinstance(item["okpd2"], str):
            raise ValueError(f"ОКПД2 для строки {row} должен быть строкой")
        code = normalize_text(item["okpd2"])
        if not code:
            raise ValueError(f"Пустой ОКПД2 для строки {row}")
        if row in result:
            raise ValueError(f"ОКПД2 для строки {row} указан повторно")
        result[row] = code
    return result
